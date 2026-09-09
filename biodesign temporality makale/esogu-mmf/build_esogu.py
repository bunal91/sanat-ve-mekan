# -*- coding: utf-8 -*-
import re, copy, os
from docx import Document
from docx.shared import Pt, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
from docx.table import Table

D = os.path.dirname(os.path.abspath(__file__))
TMPL = os.path.join(D, 'tmpl2.docx')
OUT  = os.path.join(D, 'ESOGU_MMF_Beyond_Permanence.docx')

TITLE_EN = "BEYOND PERMANENCE: REFRAMING TEMPORALITY AS A DESIGN VALUE THROUGH BIODESIGN"
TITLE_TR = "KALICILIĞIN ÖTESİNDE: BİYOTASARIM YOLUYLA BİR TASARIM DEĞERİ OLARAK ZAMANSALLIK"

KW_EN = ["Biodesign", "Temporality", "Permanence", "Material agency", "Design theory"]
KW_TR = ["Biyotasarım", "Zamansallık", "Kalıcılık", "Maddi faillik", "Tasarım kuramı"]

ABS_EN = ("Design theory has long treated permanence, meaning durability, stability and resistance "
"to change, as a proxy for design quality. The orientation is entangled with modernity's "
"ambition to render environments predictable and governable, and was institutionalised through "
"industrial standardisation and preservation-oriented maintenance. Contemporary ecological conditions "
"expose its limits: materials engineered to resist degradation often outlive their usefulness and "
"persist as ecological burdens, while ecological value depends on how matter circulates rather than on "
"how long one form endures. This article asks how biodesign, the practice in which living organisms "
"participate in material formation, destabilises the permanence paradigm, and how temporality might be "
"reconstructed as a legitimate design value. The study is conceptual, developed through an integrative "
"review following a theory synthesis design that brings together design history and sustainable design "
"theory, temporal design research, biodesign and living-materials research, and "
"process-philosophical and new materialist theory. Three findings follow. Permanence is a historically "
"and culturally situated value rather than a neutral technical criterion. Existing temporal design "
"scholarship has theorised time predominantly as an experiential and social condition, leaving the "
"material register underdeveloped; biodesign supplies that register, because growth, adaptation, "
"senescence and decomposition are constitutive of the material rather than external to it. "
"Impermanence can therefore be reframed not as material failure but as an ontological condition of "
"becoming, with permanence understood as a costly stabilisation. The article proposes paired value "
"orientations and six analytical dimensions for specifying temporal intent, and states the "
"boundary conditions under which permanence remains the correct value.")

ABS_TR = ("Tasarım kuramı uzun süredir kalıcılığı, yani dayanıklılığı, kararlılığı ve maddenin değişime "
"direncini, tasarım niteliğinin ölçütü saymıştır. Bu yönelim, modernitenin maddi ve toplumsal çevreyi "
"öngörülebilir ve yönetilebilir kılma savıyla iç içedir ve endüstriyel standartlaşma ile koruma odaklı "
"bakım pratikleri aracılığıyla kurumsallaşmıştır. Güncel ekolojik koşullar bu denklemin sınırlarını "
"açığa çıkarmaktadır: bozunmaya direnecek biçimde tasarlanan malzemeler çoğu kez işlevsel ömürlerini "
"aşarak ekolojik yük hâline gelmekte, ekolojik değer ise tek bir biçimin ne kadar sürdüğünden çok "
"maddenin nasıl dolaştığına bağlı olmaktadır. Bu çalışma, canlı organizmaların malzeme oluşumuna "
"katıldığı bir pratik olan biyotasarımın kalıcılık paradigmasını nasıl sarstığını ve zamansallığın "
"meşru bir tasarım değeri olarak nasıl yeniden kurulabileceğini sormaktadır. Çalışma kavramsal "
"niteliktedir; kuram sentezi tasarımıyla yürütülen bütünleştirici bir alanyazın taramasına dayanmakta, "
"tasarım tarihi ve sürdürülebilir tasarım kuramı, zamansallık odaklı tasarım araştırmaları, "
"biyotasarım ve canlı malzeme araştırmaları ile süreç felsefesi ve yeni maddeci kuramı bir araya "
"getirmektedir. Üç bulguya ulaşılmaktadır. Kalıcılık, yansız bir teknik ölçüt değil, tarihsel ve "
"kültürel olarak konumlanmış bir değerdir. Mevcut zamansallık yazını zamanı ağırlıkla deneyimsel ve "
"toplumsal bir koşul olarak kuramsallaştırmış, maddi kayıt görece gelişmemiş kalmıştır; biyotasarım "
"büyüme, uyarlanma, yaşlanma ve ayrışmayı malzemenin kurucu nitelikleri hâline getirdiği için tam da "
"bu kaydı sağlamaktadır. Geçicilik böylece maddi başarısızlık değil, bir oluş koşulu olarak "
"okunabilir; kalıcılık ise bedeli olan bir sabitlenme olarak yeniden tanımlanır. Çalışma, eşleştirilmiş "
"değer yönelimleri ile zamansal niyetin belirlenmesini sağlayan altı çözümleyici boyut önermekte ve "
"kalıcılığın doğru değer olmayı sürdürdüğü sınır koşullarını açıkça belirtmektedir.")

# ----------------------------------------------------------------- helpers
def set_cell(cell, texts, style=None, bold=False):
    """Replace a cell's paragraphs with `texts` (list of str)."""
    for p in list(cell.paragraphs)[1:]:
        p._element.getparent().remove(p._element)
    first = cell.paragraphs[0]
    for r in list(first.runs):
        r._element.getparent().remove(r._element)
    for i, t in enumerate(texts):
        p = first if i == 0 else cell.add_paragraph()
        if style: p.style = style
        run = p.add_run(t)
        run.bold = bold

def set_title(par, text):
    for r in list(par.runs):
        r._element.getparent().remove(r._element)
    par.add_run(text)

def uniq_cells(row):
    seen, out = set(), []
    for c in row.cells:
        if id(c._tc) not in seen:
            seen.add(id(c._tc)); out.append(c)
    return out

# ----------------------------------------------------------------- load
doc = Document(TMPL)
body = doc.element.body

# ---- front matter: English first (article language), then Turkish --------
set_title(doc.paragraphs[0], TITLE_EN)

t_en = Table(body[7], doc)          # was the Turkish table -> now English
c = uniq_cells(t_en.rows[0])
set_cell(c[0], ["Keywords"], bold=True); set_cell(c[1], ["Abstract"], bold=True)
c = uniq_cells(t_en.rows[1])
set_cell(c[0], KW_EN, style='Özet'); set_cell(c[1], [ABS_EN], style='Özet')

set_title(doc.paragraphs[[i for i,p in enumerate(doc.paragraphs) if p.text.startswith('MAKALENİN İNGİLİZCE')][0]], TITLE_TR)

t_tr = Table(body[11], doc)         # was the English table -> now Turkish
c = uniq_cells(t_tr.rows[0])
set_cell(c[0], ["Anahtar Kelimeler"], bold=True); set_cell(c[1], ["Öz"], bold=True)
c = uniq_cells(t_tr.rows[1])
set_cell(c[0], KW_TR, style='Özet'); set_cell(c[1], [ABS_TR], style='Özet')

# ---- clear the template's demo body (keep 0..12 and the final sectPr) ----
for el in list(body)[13:]:
    if el.tag == qn('w:sectPr'):
        continue
    body.remove(el)

for _p in doc.paragraphs:
    if _p.style.name == 'Başlık-1' and not _p.text.strip():
        _p.style = 'Normal'

final_sect = body[-1]

def add(el):
    final_sect.addprevious(el)

def para(text='', style='Paragraf', align=None, italic_spans=None):
    p = doc.add_paragraph(); p.style = style
    if align is not None: p.alignment = align
    if italic_spans:
        for chunk, it in italic_spans:
            r = p.add_run(chunk); r.italic = it
    elif text:
        p.add_run(text)
    body.remove(p._element); add(p._element)
    return p

MD_I = re.compile(r'(\*\*.+?\*\*|<i>.+?</i>)')
def spans(text):
    out=[]
    for part in MD_I.split(text):
        if not part: continue
        if part.startswith('**'): out.append((part[2:-2], True))
        elif part.startswith('<i>'): out.append((part[3:-4], True))
        else: out.append((part, False))
    return out

def set_cols(par, n):
    """End a section at `par` with n columns (continuous)."""
    sect = copy.deepcopy(final_sect)
    for tag in ('w:headerReference','w:footerReference','w:titlePg'):
        for e in sect.findall(qn(tag)): sect.remove(e)
    t = sect.find(qn('w:type'))
    if t is None:
        t = OxmlElement('w:type'); sect.insert(0, t)
    t.set(qn('w:val'), 'continuous')
    cols = sect.find(qn('w:cols'))
    if cols is None:
        cols = OxmlElement('w:cols'); sect.append(cols)
    if n == 1:
        if cols.get(qn('w:num')): del cols.attrib[qn('w:num')]
    else:
        cols.set(qn('w:num'), str(n))
    pPr = par._element.get_or_add_pPr()
    pPr.append(sect)

def add_table(rows, widths_in):
    tbl = doc.add_table(rows=len(rows), cols=len(rows[0]))
    tbl.style = 'Table Grid'
    tbl.autofit = False
    for ri, row in enumerate(rows):
        for ci, val in enumerate(row):
            cell = tbl.cell(ri, ci); cell.width = Inches(widths_in[ci])
            p = cell.paragraphs[0]; p.style = 'Normal'
            p.paragraph_format.space_before = Pt(0); p.paragraph_format.space_after = Pt(0)
            for chunk, it in spans(val):
                r = p.add_run(chunk); r.italic = it
                if ri == 0: r.bold = True
    body.remove(tbl._element); add(tbl._element)
    return tbl

# ---------------------------------------------------------------- content
src = open(os.path.join(D, 'esogu_body_raw.md'), encoding='utf-8').read()
lines = src.split('\n')
i = 0
pending_table = None
while i < len(lines):
    ln = lines[i].rstrip()
    s = ln.strip()
    if not s:
        i += 1; continue

    # table caption -> hold until the table itself is emitted (caption goes above)
    m_cap = re.match(r'^\*\*(Table \d+\..*?)\*\*$', s)
    if m_cap:
        pending_table = m_cap.group(1)
        i += 1; continue

    # markdown table
    if s.startswith('|') and i+1 < len(lines) and re.match(r'^\|[\s:|-]+\|$', lines[i+1].strip()):
        rows = [[c.strip() for c in s.strip('|').split('|')]]
        i += 2
        while i < len(lines) and lines[i].strip().startswith('|'):
            rows.append([c.strip() for c in lines[i].strip().strip('|').split('|')])
            i += 1
        ncol = len(rows[0])
        widths = [7.0/ncol]*ncol if ncol != 3 else [1.9, 1.9, 3.2]
        set_cols(para(''), 2)          # close the 2-column flow
        if pending_table:
            para('', 'Paragraf', WD_ALIGN_PARAGRAPH.CENTER, [(pending_table, False)])
            for _r in doc.paragraphs[-1].runs: _r.bold = True
            pending_table = None
        add_table(rows, widths)
        set_cols(para(''), 1)          # close the full-width island
        continue

    if s.startswith('[[FIGURE:'):
        img = os.path.join(D, s[len('[[FIGURE:'):-2].strip())
        set_cols(para(''), 2)
        p = doc.add_paragraph(); p.style = 'Normal'
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.add_run().add_picture(img, width=Inches(6.6))
        body.remove(p._element); add(p._element)
        i += 1
        # caption belongs to the same full-width island
        while i < len(lines) and not lines[i].strip():
            i += 1
        cap = lines[i].strip()
        para('', 'Paragraf', WD_ALIGN_PARAGRAPH.CENTER, spans(cap))
        set_cols(para(''), 1)
        i += 1
        continue

    if s.startswith('#'):
        lvl = len(s) - len(s.lstrip('#'))
        para(s.lstrip('#').strip(), 'Başlık-1')
        i += 1; continue

    para('', 'Paragraf', WD_ALIGN_PARAGRAPH.JUSTIFY, spans(s))
    i += 1

# ---- references ---------------------------------------------------------
para('References', 'Başlık-1')
for ln in open(os.path.join(D, 'esogu_refs.md'), encoding='utf-8'):
    ln = ln.strip()
    if ln:
        para('', 'Kaynaklar', None, spans(ln))

doc.save(OUT)
print('saved', OUT)
print('title EN chars:', len(TITLE_EN), '| title TR chars:', len(TITLE_TR))
print('abstract EN words:', len(ABS_EN.split()), '| öz TR words:', len(ABS_TR.split()))
