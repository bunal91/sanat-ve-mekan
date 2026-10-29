import zipfile, re, copy
import xml.etree.ElementTree as ET
from pptx import Presentation
from pptx.util import Pt

W='{http://schemas.openxmlformats.org/wordprocessingml/2006/main}'

# ---- read the author's own script, runs and all ----
z=zipfile.ZipFile('script.docx')
root=ET.fromstring(z.read('word/document.xml'))
paras=[]
for p in root.iter(W+'p'):
    runs=[]
    for r in p.iter(W+'r'):
        t=''.join(n.text or '' for n in r.iter(W+'t'))
        if not t: continue
        rpr=r.find(W+'rPr')
        b = rpr is not None and rpr.find(W+'b') is not None
        i = rpr is not None and rpr.find(W+'i') is not None
        if runs and runs[-1][1]==b and runs[-1][2]==i:
            runs[-1][0]+=t          # Word fragments runs for spell-check; put them back together
        else:
            runs.append([t,b,i])
    paras.append((''.join(x[0] for x in runs).strip(), runs))

heads=[(n,t) for n,(t,_) in enumerate(paras) if re.match(r'^\d+ · ', t)]
qa_start=next(n for n,(t,_) in enumerate(paras) if t=='Soru–cevap')

sections={}
for k,(n,title) in enumerate(heads):
    end = heads[k+1][0] if k+1 < len(heads) else qa_start
    num = int(title.split(' · ')[0])
    body=[]
    for t,runs in paras[n+1:end]:
        if t: body.append(runs)
    sections[num]=(title, body)

qa=[runs for t,runs in paras[qa_start+1:] if t]

# ---- write them into the deck's notes ----
prs=Presentation('deck.pptx')
NS=len(prs.slides)

from copy import deepcopy
def body_sp(ns):
    tf = ns.notes_text_frame
    return None if tf is None else tf._txBody.getparent()

TEMPLATE = None
def put(slide, header, blocks):
    global TEMPLATE
    ns = slide.notes_slide
    if ns.notes_text_frame is None and TEMPLATE is not None:
        ns.shapes._spTree.append(deepcopy(TEMPLATE))
    tf = ns.notes_text_frame
    tf.clear()
    first=True
    def para():
        nonlocal first
        if first:
            first=False
            return tf.paragraphs[0]
        return tf.add_paragraph()
    p=para()
    r=p.add_run(); r.text=header
    r.font.bold=True; r.font.size=Pt(17); r.font.name="Calibri"
    p.space_after=Pt(10)
    for runs in blocks:
        p=para(); p.space_after=Pt(9)
        for t,b,i in runs:
            r=p.add_run(); r.text=t
            r.font.bold=bool(b); r.font.italic=bool(i)
            r.font.size=Pt(16); r.font.name="Calibri"

TEMPLATE = body_sp(prs.slides[0].notes_slide)

for idx, slide in enumerate(prs.slides, start=1):
    if idx in sections:
        title, body = sections[idx]
        put(slide, title.upper(), body)
    elif idx == NS:
        put(slide, "SORU – CEVAP", [[["Thank you.", True, False]]] + qa)
    else:
        put(slide, "SLAYT %d" % idx, [])

prs.save('/home/user/sanat-ve-mekan/konferans-bildirisi/06-sunum/ICSMMEP26_Presentation_Notlu.pptx')
print("yazildi | slayt:", NS, "| bolum:", len(sections), "| soru-cevap paragrafi:", len(qa))
